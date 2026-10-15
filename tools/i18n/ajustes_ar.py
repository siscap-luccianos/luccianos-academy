#!/usr/bin/env python3
"""Ajustes de criterio sobre el Excel de Dariya (árabe marroquí) devuelto por la revisión.

Uso: python3 tools/i18n/ajustes_ar.py "entrada.xlsx" "salida.xlsx" ["cambios.csv"]

Mismas decisiones que en inglés (2026-10-14), adaptadas:
- Examen (no "evaluación") y Local/Tienda (no "sucursal/filial").
- Los ejemplos de dinero pasan a dirhams (el español sigue en pesos).
- Las terminaciones de plural del español ({n} pegado a "local", "incidencia"…) se sacan:
  el árabe no las tiene y el motor las deja vacías.
- "Te {0}: {1}" ({0} era el verbo falta/faltan en español) se reescribe.
- Flechas → y ←: se invierten (en RTL "adelante" apunta a la izquierda).
- Textos que el código arma por partes: meses con los nombres marroquíes que devuelve el navegador.
TODO lo traducido está pendiente de revisión por una persona nativa.
"""
import csv, re, sys
from openpyxl import load_workbook

COLS = {"Pantallas": (2, [4]), "Lecciones": (6, [7]), "Tareas y otros": (3, [4]), "Examen": (3, [4, 6, 8, 10])}

def terminos(es, t):
    if "valuaci" in es:  # Evaluación(es): تقييم → امتحان
        t = t.replace("التقييمات", "الامتحانات").replace("تقييمات", "امتحانات").replace("التقييم", "الامتحان").replace("تقييم", "امتحان")
    if "ucursal" in es:  # Sucursal(es): فرع → محل
        t = t.replace("الفروع", "المحلات").replace("فروع", "محلات").replace("الفرع", "المحل").replace("فرع", "محل")
    return t

DINERO = [
    ("$1.000.000 ÷ 500 تيكي = $2.000 معدل التيكي", "5000 درهم ÷ 500 تيكي = 10 دراهم معدل التيكي"),
    ("$1.000.000/500 تيكي ($2.000 كل وحدة) مقابل $1.000.000/250 تيكي ($4.000 كل وحدة)", "5000 درهم/500 تيكي (10 دراهم كل وحدة) مقابل 5000 درهم/250 تيكي (20 درهم كل وحدة)"),
    ("الثمن $5.000، التكلفة $1.500", "الثمن 50 درهم، التكلفة 15 درهم"),
]

# {n} pegado a estas palabras = terminación de plural en español
RE_PLURAL = re.compile(r"(?:local|incidencia|grave|producto|acotado|depende|excluido)\{(\d+)\}")

FRASES = {
    # {0} era el verbo (falta/faltan): en dariya no hace falta
    "Para participar necesitás aprobar el examen de todos tus módulos. Te {0}: {1}.": "باش تشارك خاصك تنجح ف الامتحان ديال جميع الوحدات ديالك. باقي خاصك: {1}.",
    "Para participar del Desafío Diario necesitás aprobar el examen de todos tus módulos. Te {0}: {1}.": "باش تشارك ف التحدي ديال النهار خاصك تنجح ف الامتحان ديال جميع الوحدات ديالك. باقي خاصك: {1}.",
    "Hay {0} {1} que depende{2} de ella:": "كاين {0} {1} مرتابطين بيها:",
}

MESES = ["يناير", "فبراير", "مارس", "أبريل", "ماي", "يونيو", "يوليوز", "غشت", "شتنبر", "أكتوبر", "نونبر", "دجنبر"]

EXTRA_UI = {
    "versiones": "نسخ",
    # Tarjeta "Idiomas — vista previa" de Configuración (Admin)
    "Idiomas — vista previa": "اللغات — معاينة",
    "Los idiomas están abiertos: el inglés lo puede elegir cualquier persona en Mi perfil → Idioma, y el dariya solo los locales de Marruecos.": "اللغات مفتوحة: الإنجليزية يقدر يختارها أي واحد ف البروفايل ديالي ← اللغة، والدارجة غير للمحلات ديال المغرب.",
    "Por ahora el inglés y el dariya los ven solo vos y las personas que elijas acá (lo eligen en Mi perfil → Idioma). Para el resto la app sigue igual, en español.": "دابا الإنجليزية والدارجة كيشوفهم غير نتا والناس اللي كتختار هنا (كيختارو اللغة ف البروفايل ديالي ← اللغة). الباقي كيبقى التطبيق كيف هو، بالإسبانية.",
    "Listo. Esas personas ya pueden elegir el idioma en Mi perfil (si tenían la app abierta, que la recarguen).": "تمّ. هاد الناس دابا يقدرو يختارو اللغة ف البروفايل ديالي (إلا كان التطبيق مفتوح عندهم، خاصهم يعاودو يحمّلوه).",
}
for i, mes in enumerate(MESES):
    EXTRA_UI["{0} de " + mes] = "{0} " + mes

FLECHAS = str.maketrans({"→": "←", "←": "→"})   # en derecha a izquierda "adelante" apunta a la izquierda

def main():
    if len(sys.argv) < 3: sys.exit(__doc__)
    wb = load_workbook(sys.argv[1])
    log = []
    for hoja, (col_es, cols_t) in COLS.items():
        ws = wb[hoja]
        for r in range(2, ws.max_row + 1):
            for ct in cols_t:
                ce = ct - 1 if hoja == "Examen" else col_es
                es, t = ws.cell(r, ce).value, ws.cell(r, ct).value
                if not es or not t: continue
                es, t = str(es), str(t)
                nuevo = terminos(es, t)
                for a, b in DINERO: nuevo = nuevo.replace(a, b)
                if hoja == "Pantallas":
                    if es in FRASES: nuevo = FRASES[es]
                    else:
                        for m in RE_PLURAL.finditer(es):
                            nuevo = nuevo.replace("{%s}" % m.group(1), "")
                if hoja == "Pantallas" and ("→" in es or "←" in es):
                    nuevo = nuevo.translate(FLECHAS)
                if nuevo != t:
                    ws.cell(r, ct).value = nuevo
                    log.append([hoja, ws.cell(r, 1).value, es[:70].replace("\n", " / "), t[:100].replace("\n", " / "), nuevo[:100].replace("\n", " / ")])
    ws = wb["Pantallas"]
    existentes = {str(ws.cell(r, 2).value) for r in range(2, ws.max_row + 1)}
    for i, (es, ar) in enumerate(EXTRA_UI.items(), 1):
        if es not in existentes:
            ws.append([f"UI-X{i:03d}", es, "(parte de una frase / pantalla de Admin)", ar, "Agregado en el ajuste — revisar"])
            log.append(["Pantallas", f"UI-X{i:03d}", es, "(no existía)", ar])
    wb.save(sys.argv[2])
    if len(sys.argv) > 3:
        with open(sys.argv[3], "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f); w.writerow(["Hoja", "ID", "Español", "Antes", "Después"]); w.writerows(log)
    print(f"{len(log)} traducciones ajustadas → {sys.argv[2]}")

main()
