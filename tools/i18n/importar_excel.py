#!/usr/bin/env python3
"""Importa el Excel traducido y genera los diccionarios de la app.

Uso (desde la raíz del repo):
    python3 tools/i18n/importar_excel.py "traducido.xlsx" [--salida js/i18n] [--codigo en] [--estricto]

Genera <salida>/<codigo>.js (interfaz) y <salida>/<codigo>.contenido.js
(lecciones, preguntas, tareas…). Lo que no esté traducido se omite: la app lo
muestra en español. Sin --estricto, los problemas se informan pero igual se
genera; con --estricto no se escribe nada si hay algún problema.
"""
import json, re, sys, os, collections
from openpyxl import load_workbook

IDIOMA_NOMBRE = {"en": "inglés (EE. UU.)"}

def marcadores(s): return sorted(re.findall(r"\{\d+\}", s))

def main():
    a = sys.argv[1:]
    if not a: sys.exit(__doc__)
    ruta = a[0]
    salida = a[a.index("--salida")+1] if "--salida" in a else "js/i18n"
    codigo = a[a.index("--codigo")+1] if "--codigo" in a else "en"
    estricto = "--estricto" in a
    wb = load_workbook(ruta, data_only=True)

    problemas, avisos = [], []
    ui, cont = {}, {}
    estad = collections.Counter()

    def agregar(dic, es, tr, donde, multilinea):
        es = es if isinstance(es, str) else ("" if es is None else str(es))
        tr = tr if isinstance(tr, str) else ("" if tr is None else str(tr))
        if not es.strip(): return
        estad[donde + " total"] += 1
        if not tr.strip():
            estad[donde + " sin traducir"] += 1; return
        me, mt = marcadores(es), marcadores(tr)
        sobran = [m for m in mt if m not in me]
        faltan = [m for m in me if m not in mt]
        if sobran or (len(mt) > len(set(mt)) and mt != me):
            problemas.append(f"[{donde}] ERROR marcadores que no existen en el español {sobran}: «{es[:60]}» → «{tr[:60]}»")
        elif faltan:
            avisos.append(f"[{donde}] falta el marcador {faltan} (válido solo si es una terminación/plural que ese idioma no tiene): «{es[:60]}» → «{tr[:60]}»")
        if multilinea and es.count("\n") != tr.count("\n"):
            problemas.append(f"[{donde}] distinta cantidad de líneas ({es.count(chr(10))+1} vs {tr.count(chr(10))+1}): «{es[:60]}»")
        if es in dic and dic[es] != tr:
            problemas.append(f"[{donde}] mismo español con dos traducciones: «{es[:60]}» → «{dic[es][:40]}» / «{tr[:40]}» (se queda la primera)")
            return
        dic[es] = tr

    def col(ws, *nombres):
        enc = [str(c.value or "").strip() for c in ws[1]]
        for n in nombres:
            for i, e in enumerate(enc):
                if e.startswith(n): return i
        sys.exit(f"Hoja «{ws.title}»: no encuentro la columna «{nombres[0]}». ¿Se modificó la fila de encabezados?")

    ws = wb["Pantallas"]; ie, it = col(ws, "Español"), col(ws, "Traducci")
    for r in ws.iter_rows(min_row=2, values_only=True): agregar(ui, r[ie], r[it], "Pantallas", False)

    ws = wb["Lecciones"]; ie, it = col(ws, "Español"), col(ws, "Traducci")
    for r in ws.iter_rows(min_row=2, values_only=True): agregar(cont, r[ie], r[it], "Lecciones", True)

    ws = wb["Examen"]
    enc = [str(c.value or "").strip() for c in ws[1]]
    pares = [(2, 3), (4, 5), (6, 7), (8, 9)]
    if not (enc[2].startswith("Pregunta (ES") and enc[3].startswith("Pregunta (trad") and enc[4].startswith("Opción 1 (ES")):
        sys.exit("Hoja «Examen»: cambió el orden de las columnas.")
    for r in ws.iter_rows(min_row=2, values_only=True):
        for e, t in pares: agregar(cont, r[e], r[t], "Examen", False)

    ws = wb["Tareas y otros"]; ie, it = col(ws, "Español"), col(ws, "Traducci")
    for r in ws.iter_rows(min_row=2, values_only=True): agregar(cont, r[ie], r[it], "Tareas y otros", False)

    print("Resumen:")
    for k in sorted(estad): print(f"  {k}: {estad[k]}")
    print(f"  traducidos → interfaz: {len(ui)} · contenido: {len(cont)}")
    if avisos:
        print(f"\n{len(avisos)} aviso(s) (no bloquean):")
        for p_ in avisos[:30]: print("  -", p_)
    if problemas:
        print(f"\n{len(problemas)} problema(s):")
        for p in problemas[:60]: print("  -", p)
        if len(problemas) > 60: print(f"  … y {len(problemas)-60} más")
        if estricto: sys.exit("\n--estricto: no se escribió nada.")

    nombre = IDIOMA_NOMBRE.get(codigo, codigo)
    cab_ui = f'''/* ============================
   Lucciano's Academy
   i18n/{codigo}.js — Diccionario de la interfaz: español → {nombre}

   Un texto por línea: "texto en español": "traducción".
   {{0}}, {{1}}… son las partes variables (nombres, números). Para revisar o
   corregir una traducción alcanza con editar el texto de la derecha.
   Un texto que no está acá se muestra en español.
   GENERADO por tools/i18n/importar_excel.py desde el Excel de traducción.
=============================*/

'''
    cab_c = "// Contenido de planillas (lecciones, preguntas, tareas…) traducido — " + nombre + ".\n// Clave: texto exacto en español tal como está en la planilla. GENERADO por tools/i18n/importar_excel.py.\n"
    def js(dic, cab):
        cuerpo = "export default {\n" + "".join(f"    {json.dumps(k, ensure_ascii=False)}: {json.dumps(v, ensure_ascii=False)},\n" for k, v in dic.items()) + "};\n"
        return cab + cuerpo
    # Seguro: no pisar un diccionario ya cargado con uno mucho más chico
    # (ej. importar un Excel a medio traducir) salvo que se pida con --forzar.
    for nom, dic in ((f"{codigo}.js", ui), (f"{codigo}.contenido.js", cont)):
        destino = os.path.join(salida, nom)
        if os.path.exists(destino) and "--forzar" not in a:
            previo = len(re.findall(r'^    "', open(destino, encoding="utf8").read(), re.M))
            if previo and len(dic) < previo * 0.8:
                sys.exit(f"\nNo se escribió nada: {nom} tiene {previo} textos y el Excel traería {len(dic)}. Si es lo que querés, agregá --forzar.")
    os.makedirs(salida, exist_ok=True)
    open(os.path.join(salida, f"{codigo}.js"), "w", encoding="utf8").write(js(ui, cab_ui))
    open(os.path.join(salida, f"{codigo}.contenido.js"), "w", encoding="utf8").write(js(cont, cab_c))
    print(f"\nOK → {salida}/{codigo}.js y {salida}/{codigo}.contenido.js")

main()
