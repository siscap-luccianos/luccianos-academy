#!/usr/bin/env python3
"""Exporta TODO el texto en español de la app a un Excel para traducir.

Uso (desde la raíz del repo):
    python3 tools/i18n/exportar_excel.py "salida.xlsx" [--idioma "English (US)"] [--con-borrador]

--con-borrador rellena la columna "Traducción" con la traducción que ya está
cargada en js/i18n (sirve para revisar o para probar la ida y vuelta).
Los datos de lecciones/preguntas salen de tools/i18n/datos/*.json (volcado de
la planilla de producción); la interfaz sale de las claves de js/i18n/en.js.
"""
import json, re, subprocess, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DATOS = os.path.join(RAIZ, "tools", "i18n", "datos")
CURSOS = {'1':'Cafetería','2':'Heladería','3':'Icepops','4':'Pastelería','5':'Chocolatería','6':'Sistema y Caja','7':'Atención al Cliente','8':'Responsables de Local y Turno'}

def cargar_js(ruta):
    """Lee un `export default {...}` con node y lo devuelve como dict."""
    cod = "import(%s).then(m=>process.stdout.write(JSON.stringify(m.default)))" % json.dumps("file://" + ruta)
    return json.loads(subprocess.check_output(["node", "-e", cod]))

def norm(k): return re.sub(r"\s+", " ", k).strip()

def main():
    args = sys.argv[1:]
    if not args: sys.exit(__doc__)
    salida = args[0]
    idioma = args[args.index("--idioma")+1] if "--idioma" in args else "English (US)"
    borrador = "--con-borrador" in args

    ui = cargar_js(os.path.join(RAIZ, "js/i18n/en.js"))
    cont = cargar_js(os.path.join(RAIZ, "js/i18n/en.contenido.js"))
    arch = {}
    for f, ks in json.load(open(os.path.join(DATOS, "ui_archivos_raw.json"))).items():
        for k in ks: arch.setdefault(norm(k), f)
    lec = json.load(open(os.path.join(DATOS, "lecciones.json")))["values"]
    ev = json.load(open(os.path.join(DATOS, "evaluaciones.json")))["values"]
    otros = json.load(open(os.path.join(DATOS, "tareas_y_otros.json")))

    def trad(es):  # borrador existente (o vacío)
        return (cont.get(es.strip()) or cont.get(es) or ui.get(es) or "") if borrador else ""

    wb = Workbook()
    H = PatternFill("solid", start_color="1F3A5F"); HF = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    B = Font(name="Arial", size=10); BLOQ = PatternFill("solid", start_color="F2F4F7"); AMAR = PatternFill("solid", start_color="FFF8DC")
    WRAP = Alignment(wrap_text=True, vertical="top")

    def hoja(nombre, cols, anchos, filas, col_es, col_tr):
        ws = wb.create_sheet(nombre)
        ws.append(cols)
        for i in range(1, len(cols)+1):
            c = ws.cell(1, i); c.fill = H; c.font = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = anchos[i-1]
        for f in filas: ws.append(f)
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.font = B; c.alignment = WRAP
                if c.column != col_tr: c.fill = BLOQ
                else: c.fill = AMAR
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        return ws

    # ── Instrucciones
    ws = wb.active; ws.title = "Instrucciones"
    texto = INSTRUCCIONES.replace("{IDIOMA}", idioma)
    for linea in texto.split("\n"):
        ws.append([linea])
        c = ws.cell(ws.max_row, 1); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Arial", size=11 if linea.isupper() or linea.endswith(":") else 10, bold=linea.isupper() or linea.endswith(":"))
    ws.column_dimensions["A"].width = 140

    # ── Pantallas (interfaz)
    filas = []
    for i, es in enumerate(ui.keys(), 1):
        filas.append([f"UI-{i:04d}", es, arch.get(norm(es), ""), trad(es), ""])
    hoja("Pantallas", ["ID", "Español (NO modificar)", "Dónde aparece", "Traducción", "Notas"], [10, 70, 36, 70, 30], filas, 2, 4)

    # ── Lecciones (todas las columnas de texto, todas las lecciones)
    h = lec[0]; ix = {n: i for i, n in enumerate(h)}
    campos = [("titulo","Título"),("objetivo","Objetivo"),("procedimiento","Procedimiento"),("errores","Errores frecuentes"),("buenasPracticas","Buenas prácticas"),("consejo","Consejo"),("resumen","Resumen"),("manualLabel","Texto del enlace al manual")]
    vistos = {}; filas = []
    for r in lec[1:]:
        r = r + [""]*(len(h)-len(r))
        curso = CURSOS.get(r[ix["cursoId"]], r[ix["cursoId"]])
        vis = "Sí" if not r[ix["aplicaA"]].strip() else "Solo " + r[ix["aplicaA"]].strip()
        for c, etq in campos:
            t = r[ix[c]].strip()
            if not t: continue
            if t in vistos: vistos[t][0] += 1; continue
            fila = [f"LEC-{len(vistos)+1:04d}", curso, int(r[ix["orden"]] or 0), r[ix["titulo"]], etq, t, trad(t), vis, ""]
            vistos[t] = [1, fila]; filas.append(fila)
    ordenc = list(CURSOS.values())
    filas.sort(key=lambda f: (ordenc.index(f[1]) if f[1] in ordenc else 99, f[2]))
    hoja("Lecciones", ["ID", "Curso", "Nº", "Lección", "Campo", "Español (NO modificar)", "Traducción", "Se muestra en EE. UU.", "Notas"], [10, 22, 6, 30, 18, 80, 80, 18, 30], filas, 6, 7)

    # ── Examen (una fila por pregunta, con sus 3 opciones)
    eh = ev[0]; ex = {n: i for i, n in enumerate(eh)}
    filas = []; vistas = {}
    for r in ev[1:]:
        if not r or not r[0]: continue
        r = r + [""]*(12-len(r))
        vis = "Sí" if not r[10].strip() else "Solo " + r[10].strip()
        corr = r[7]
        filas.append([f"EXA-{len(filas)+1:04d}", CURSOS.get(r[1], r[1]), r[2], trad(r[2]), r[4], trad(r[4]), r[5], trad(r[5]), r[6], trad(r[6]), corr, vis, ""])
    ordenc = list(CURSOS.values())
    filas.sort(key=lambda f: (ordenc.index(f[1]) if f[1] in ordenc else 99, f[0]))
    hoja("Examen", ["ID", "Curso", "Pregunta (ES - NO modificar)", "Pregunta (traducción)", "Opción 1 (ES)", "Opción 1 (traducción)", "Opción 2 (ES)", "Opción 2 (traducción)", "Opción 3 (ES)", "Opción 3 (traducción)", "Opción correcta (nº)", "Se muestra en EE. UU.", "Notas"],
         [10, 20, 45, 45, 28, 28, 28, 28, 28, 28, 10, 16, 24], filas, 3, 4)
    # las columnas de traducción del examen son 4, 6, 8, 10: se pintan aparte
    ws = wb["Examen"]
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.fill = AMAR if c.column in (4, 6, 8, 10) else BLOQ

    # ── Tareas y otros
    filas = []
    for tipo, clave in [("Tarea de Gestión", "tareas"), ("Detalle de tarea", "detalles"), ("Sub-ítem de tarea", "subitems"), ("Título de manual", "manuales"), ("Recurso", "recursos"), ("Curso", "cursos")]:
        for es in otros[clave]:
            filas.append([f"OTR-{len(filas)+1:04d}", tipo, es, trad(es), ""])
    hoja("Tareas y otros", ["ID", "Tipo", "Español (NO modificar)", "Traducción", "Notas"], [10, 22, 60, 60, 30], filas, 3, 4)

    wb.save(salida)
    print("OK →", salida)

INSTRUCCIONES = """INSTRUCCIONES PARA TRADUCIR (Lucciano's Academy)
Idioma de destino: {IDIOMA}
 
QUÉ ES ESTE ARCHIVO:
Contiene todo el texto en español de la app Lucciano's Academy (capacitación de personal de una cadena de heladerías, cafeterías, pastelería y chocolatería). Hay 4 hojas: Pantallas (botones, menús, mensajes), Lecciones, Examen y Tareas y otros.
 
QUÉ HAY QUE HACER:
Completar la columna amarilla "Traducción" (en Examen, las 4 columnas amarillas de traducción). No modificar las columnas grises, ni agregar, borrar ni reordenar filas.
Cuando termines, devolvé el mismo archivo .xlsx: se carga a la app con un script que lee esas columnas.
 
REGLAS (importantes, el script las controla):
1. Marcadores {0}, {1}, {2}…: son partes variables (nombres, números). Van TAL CUAL en la traducción, los mismos y en cualquier posición que corresponda al idioma. No se traducen ni se renumeran.
2. Saltos de línea y estructura: en Lecciones, respetar el mismo número de líneas, los números "1) 2) 3)" y las viñetas con sangría "   - ". Una línea del español = una línea de la traducción.
3. Mismo texto en español = misma traducción en todo el archivo.
4. Nombres de marca y de producto NO se traducen: Lucciano's, Avella, Geladot, Geladot's, Conitos, Squares, Viennesi, Pistacchio, Crock, Oli, Vito, Enzo, Nespresso y similares.
5. Medidas: se mantienen las métricas (g, ml, cm). Las temperaturas siguen en °C; agregar entre paréntesis el equivalente en °F si el país lo usa.
6. Tono: cercano y directo, de equipo a equipo (el español usa "vos"). Traducir como lo diría alguien del rubro en el país de destino, no palabra por palabra.
7. Botones y menús: textos cortos, parecidos en largo al original (hay poco espacio en el celular).
8. Si un texto no se entiende o depende de algo específico de Argentina (Mercado Pago, Factura A, IVA, PedidosYa…), traducirlo igual y anotar la duda en la columna "Notas". No dejar filas en blanco salvo que sea imposible.
9. Las filas con "Solo Argentina/Chile" en "Se muestra en EE. UU." se pueden traducir igual (opcional).
 
PROMPT SUGERIDO PARA CLAUDE (pegarlo junto con el archivo):
You are translating the interface and training content of "Lucciano's Academy", a staff-training app for a chain of ice cream, gelato, café, pastry and chocolate shops, from Argentine Spanish into {IDIOMA}. Fill in the "Traducción" column(s) of every row in every sheet, following the rules in the "Instrucciones" sheet exactly (keep {0}/{1} placeholders, keep line breaks and "1) ... - ..." structure, do not translate brand/product names, same Spanish text = same translation, keep metric units). Use the vocabulary of ice cream / gelato / café / chocolate shops in the target country. Do not change any other column. Work sheet by sheet in batches of about 100-150 rows so nothing is skipped, and return the complete .xlsx at the end.
 
CÓMO SE CARGA DESPUÉS:
Se pasa el archivo devuelto a quien mantiene la app; un script (tools/i18n/importar_excel.py) lo revisa (marcadores, líneas, filas vacías, textos repetidos con distinta traducción) y genera los diccionarios de la app. Lo que no esté traducido sigue apareciendo en español.
"""

main()
